"""Synthetic test files in several bank formats (not real data)."""
import csv, pathlib
from openpyxl import Workbook
from datetime import datetime, timedelta
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

OUT = pathlib.Path(__file__).parent / 'samples'; OUT.mkdir(exist_ok=True)
T0 = datetime(2026, 8, 18, 11, 0, 0)
COMP, A1, A2, A3, B1, B2 = '001101234567', '41230098765', '30012345678', '50100234567890', '918020012345678', '112233445566'

def icici_complainant():
    wb = Workbook(); ws = wb.active; ws.title = 'OpTransactionHistory'
    ws.append(['ICICI BANK LTD — DETAILED STATEMENT']); ws.append(['Account Number:', COMP]); ws.append(['Account Name: TEST VICTIM JOSE']); ws.append(['IFSC: ICIC0000011']); ws.append([])
    ws.append(['Tran_ID', 'Tran_Date', 'Narration', 'Dr_Amt', 'Cr_Amt', 'Balance', 'pstd_dt'])
    bal = 900000.0; rows = [
        ('S1001', T0 - timedelta(days=2), 'NEFT/SALARY/ACME LTD', 0, 85000),
        ('S1002', T0, f'IMPS/P2A/623011223344/RAVI K/BARB/X{A1[-4:]}', 250000, 0),
        ('S1003', T0 + timedelta(minutes=9), f'IMPS/P2A/623011223355/SUNIL/SBIN/{A2}', 180000, 0),
        ('S1004', T0 + timedelta(minutes=21), 'UPI/DR/623011229999/MEENA/HDFC/meena77@okhdfcbank', 95000, 0),
        ('S1005', T0 + timedelta(hours=5), 'ATM WDL/KOCHI', 5000, 0)]
    for tid, dt, n, d, c in rows:
        bal = bal + c - d
        ws.append([tid, dt.strftime('%d/%m/%Y'), n, d or None, c or None, round(bal, 2), dt.strftime('%d/%m/%Y %I:%M:%S %p')])
    wb.save(OUT / 'ICICI_complainant.xlsx')

def bob_l1():  # newest first, date+time in one column, separate Transaction ID
    wb = Workbook(); ws = wb.active
    ws.append(['Bank of Baroda']); ws.append([f'A/c No: {A1}']); ws.append(['Name : RAVI KUMAR']); ws.append([])
    ws.append(['Transaction Date', 'Transaction ID', 'Transaction Details', 'Debit Amount', 'Credit Amount', 'Line Balance'])
    ev = [(T0 - timedelta(days=1), 'T0', 'IMPS/P2A/600000000001/ACCOUNT VALIDATION', 0, 1),
          (T0 + timedelta(minutes=1), 'T1', 'IMPS/P2A/623011223344/TEST VICTIM/ICIC', 0, 250000),
          (T0 + timedelta(minutes=35), 'T2', f'IMPS/P2A/623099887766/ANIL/UTIB/{B1}', 200000, 0),
          (T0 + timedelta(minutes=50), 'T3', 'ATM WDL/NUH', 40000, 0)]
    bal = 1200.0; out = []
    for dt, tid, n, d, c in ev:
        bal = bal + c - d; out.append([dt.strftime('%d-%m-%Y %H:%M:%S'), tid, n, d, c, round(bal, 2)])
    for r in reversed(out): ws.append(r)
    wb.save(OUT / 'BOB_L1_ravi.xlsx')

def sbi_l1_csv():  # CSV, dates without time, single Amount + Dr/Cr
    with open(OUT / 'SBI_L1_sunil.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['State Bank of India']); w.writerow(['Account No', A2]); w.writerow([])
        w.writerow(['Txn Date', 'Value Date', 'Description', 'Ref No./Cheque No.', 'Amount', 'Dr/Cr', 'Balance'])
        bal = 3000.0
        for dt, n, ref, amt, dc in [(T0, 'BY TRANSFER-IMPS/623011223355/TEST VICTIM', '623011223355', 180000, 'CR'), (T0, f'TO TRANSFER-IMPS/623055556666/{B2}/ZAKIR', '623055556666', 170000, 'DR')]:
            bal = bal + (amt if dc == 'CR' else -amt)
            w.writerow([dt.strftime('%d/%m/%Y'), dt.strftime('%d/%m/%Y'), n, ref, f'{amt:,.2f}', dc, f'{bal:,.2f}'])

def hdfc_l1_pdf():
    c = canvas.Canvas(str(OUT / 'HDFC_L1_meena.pdf'), pagesize=landscape(A4)); W, H = landscape(A4)
    c.setFont('Helvetica', 9); c.drawString(40, H - 40, 'HDFC BANK Statement of account'); c.drawString(40, H - 55, f'Account No : {A3}   IFSC : HDFC0001234'); c.drawString(40, H - 70, 'Name : MEENA S')
    cols = [(40, 'Date'), (95, 'Narration'), (330, 'Chq./Ref.No.'), (420, 'Value Dt'), (500, 'Withdrawal Amt.'), (600, 'Deposit Amt.'), (700, 'Closing Balance')]
    y = H - 100
    for x, t in cols: c.drawString(x, y, t)
    bal = 500.0
    rows = [(T0 + timedelta(minutes=22), 'UPI-TEST VICTIM-623011229999', '623011229999', 0, 95000), (T0 + timedelta(hours=1), 'UPI-ZAKIR-zakir@ybl-623077778888', '623077778888', 90000, 0)]
    for dt, n, ref, d, cr in rows:
        y -= 16; bal = bal + cr - d
        c.drawString(40, y, dt.strftime('%d/%m/%y')); c.drawString(95, y, n); c.drawString(330, y, ref); c.drawString(420, y, dt.strftime('%d/%m/%y'))
        if d: c.drawRightString(566, y, f'{d:,.2f}')
        if cr: c.drawRightString(652, y, f'{cr:,.2f}')
        c.drawRightString(765, y, f'{bal:,.2f}')
        y -= 12; c.drawString(95, y, 'continued narration line')
    c.save()

def kyc():
    wb = Workbook(); ws = wb.active
    ws.append(['Account Number', 'Account Holder Name', 'Registered Mobile', 'Alternate Mobile No', 'Email ID', 'PAN', 'IFSC', 'Address', 'UPI ID'])
    ws.append([COMP, 'TEST VICTIM JOSE', '9847000001', '', 'victim@example.com', 'ABCPJ1234K', 'ICIC0000011', 'KOCHI', ''])
    ws.append([A1, 'RAVI KUMAR', '9812000002', '7300000003 / 9999000004', '', 'BBBPK1111A', 'BARB0NUHXXX', 'NUH, HARYANA', ''])
    ws.append([A2, 'SUNIL YADAV', '9812000002', '', '', 'CCCPY2222B', 'SBIN0001234', 'JAMTARA', ''])
    ws.append([A3, 'MEENA S', '8800000005', '', '', '', 'HDFC0001234', 'DELHI', 'meena77@okhdfcbank'])
    wb.save(OUT / 'KYC_reply.xlsx')

def iplog():
    wb = Workbook(); ws = wb.active
    ws.append(['Account No', 'Login Date Time (GMT)', 'IP Address', 'Port', 'Channel', 'Device ID'])
    to_utc = lambda d: (d - timedelta(hours=5, minutes=30)).strftime('%d/%m/%Y %H:%M:%S')
    ws.append([COMP, to_utc(T0 - timedelta(days=5)), '117.216.10.20', '', 'iMobile', 'SM-A515F'])
    ws.append([COMP, to_utc(T0 - timedelta(minutes=3)), '49.36.12.34', '40211', 'iMobile', 'DEV-X'])
    ws.append([A1, to_utc(T0 + timedelta(minutes=33)), '49.36.12.34', '40300', 'MB', 'DEV-X'])
    ws.append([A2, to_utc(T0 - timedelta(minutes=5)), '100.64.3.9', '51000', 'YONO', 'DEV-Y'])
    wb.save(OUT / 'IPLOG_ICICI_BOB_SBI.xlsx')

def ipdr():
    wb = Workbook(); ws = wb.active
    ws.append(['MSISDN', 'Public IP', 'Public Port Start', 'Public Port End', 'Start Time', 'End Time', 'IMEI', 'Cell ID'])
    ws.append(['917300000003', '49.36.12.34', 40000, 40999, (T0 - timedelta(hours=1)).strftime('%d/%m/%Y %H:%M:%S'), (T0 + timedelta(hours=2)).strftime('%d/%m/%Y %H:%M:%S'), '356789012345678', '404-10-5555'])
    wb.save(OUT / 'IPDR_jio.xlsx')

def cdr():
    wb = Workbook(); ws = wb.active
    ws.append(['CDR of 7300000003']); ws.append(['Calling No', 'Called No', 'Call Date', 'Call Time', 'Duration', 'Call Type', 'First Cell ID', 'IMEI'])
    for k, (a, b, typ) in enumerate([('7300000003', '9847000001', 'MOC'), ('9847000001', '7300000003', 'MTC'), ('7300000003', '9812000002', 'MOC'), ('7300000003', '9999888877', 'MOC')]):
        dt = T0 - timedelta(minutes=30) + timedelta(minutes=12 * k)
        ws.append([a, b, dt.strftime('%d/%m/%Y'), dt.strftime('%H:%M:%S'), 60 + k, typ, '404-10-5555', '356789012345678'])
    wb.save(OUT / 'CDR_7300000003.xlsx')

for f in [icici_complainant, bob_l1, sbi_l1_csv, hdfc_l1_pdf, kyc, iplog, ipdr, cdr]: f()
print('samples in', OUT, sorted(p.name for p in OUT.iterdir()))
